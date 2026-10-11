"""Build a phone agent from a bank's brief. One command, no manual step.

    python3 builder/build.py <brief file or folder> [more brief files] [<tool API file>] [<address of the bank's services>] [--slot NAME] [--no-deploy] [--spec FILE]

The tool API file is the .json file that holds an OpenAPI description; every other file is brief material (a brief that
builds on another: pass both, base first). A folder stands for all files in it: Markdown and other text, Word, PDF, Excel.
Without a tool API file the builder writes one from the service description in the documents. An address (https://...)
is where the phone agent's tools reach the bank's services. What it does:
  1. reads the brief and the tool API file into a policy spec (model calls, see extract.py)
  2. compiles the spec into a workflow (engine.py), with no model involved
  3. checks the workflow and runs its in-call code against a fake bank
  4. saves it to Freya (the same slot updates the same workflow), exports it to submission/<slot>.workflow.json
  5. asks the agent for its first line over the chat route, to see that the workflow starts

It writes a plain-text trace to traces/<slot>.log and the spec to out/<slot>.spec.json.
--spec FILE skips step 1 and compiles that spec (for development). Needs FREYA_API_KEY, and a Claude login or ANTHROPIC_API_KEY.
"""
import json, os, pathlib, sys, time
ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "tests"))
from builder import engine, extract, llm

def selftest(spec, openapi, trace):
    """The compiled code against a fake bank, on an account made up from the tool API file. Raises if a basic call breaks."""
    from local_sim import Sim
    api = engine.read_api(openapi); P = spec["product"]
    api = dict(api, get_account=api[engine.service(api, spec.get("tools") or {}, "get_account")])        # the account service may have another name
    types = [t for t in P.get("types", {}) if t != "*"] or ["product"]
    def product(i, t):
        p = {P["id_field"]: f"p{i + 1}", P["type_field"]: t, P["full_field"]: 300.0 * (i + 1), P["days_late_field"]: (spec["bands"][min(i, len(spec["bands"]) - 1)].get("min") or 1)}
        for k in api["get_account"]["response"].get(P["list_field"], {}).get("items", {}).get("properties", {}):
            if k not in p: p[k] = "1234" if "last" in k else 90.0 * (i + 1)
        return p
    identity = {q["field"]: {"date": "1980-02-03", "digits": "1234"}.get(q["kind"], "AB12CD") for q in spec["identity"]["questions"]}     # a text answer that upper-casing and removing spaces leave as it is
    now = str(spec.get("test_variables", {}).get(spec["now_variable"]) or "2026-01-05T11:00")
    ran = []
    for n in ((1, 3) if spec["notes"].get("enabled") else (1,)):
        ps = [product(i, types[i % len(types)]) for i in range(n)]
        s = Sim(spec, ps, now, identity=identity).login()
        st = s.state
        assert st["stage"] == "verified" and len(st["products"]) == n, "the account did not load"
        for p in st["products"]: assert p["deadline"] and p["floor"] <= p["full"], "a product has no deadline or a floor above the full amount"
        d = min(p["deadline"] for p in st["products"])
        s.tool("set_promise", product_id="all" if n > 1 else st["products"][0]["id"], amount="full", date=d); s.tool("confirm")
        for _ in range(n - 1 if P.get("confirm") == "each" else 0): s.tool("confirm")        # one read-back and one yes per product
        s.tool("finish", outcome=spec["outcomes"]["promise_recorded"][0])
        assert s.ledger["outcome"] == spec["outcomes"]["promise_recorded"][0] and len(s.ledger["promises"]) == n, "a simple promise call did not end as a promise"
        assert [c["tool"] for c in s.ledger["calls"]][-1] == "close_call" and [c["tool"] for c in s.ledger["calls"]].count("close_call") == 1, "the call was not closed exactly once at the end"
        s = Sim(spec, ps, now, identity=identity).login(); s.hang_up()
        assert s.ledger["outcome"], "a call that breaks off was not closed"
        ran.append(f"{n} product(s): promise on {d}, break-off closes as {s.ledger['outcome']}/{s.ledger['reason']}")
    trace.log("SELF-TEST against a fake bank passed: " + "; ".join(ran))

def inputs(args):
    """(brief files, tool API file or None, address of the bank's services or None) from the arguments. Folders are opened;
    the tool API file is recognised by what is in it, so a code list that comes as .json stays brief material."""
    url = next((a for a in args if a.lower().startswith(("http://", "https://"))), None)
    files = []
    for a in args:
        if a == url: continue
        p = pathlib.Path(a)
        files += sorted(x for x in p.rglob("*") if x.is_file() and not x.name.startswith((".", "~$"))) if p.is_dir() else [p]
    def is_api(f):
        try: d = json.loads(f.read_text())
        except Exception: return False
        return isinstance(d, dict) and isinstance(d.get("paths"), dict)
    apis = [f for f in files if f.suffix.lower() == ".json" and is_api(f)]
    api = apis[-1] if apis else None
    return [f for f in files if f != api], api, url

LIMIT = 120000          # characters of one document that go to the model

def read_text(f):
    """A brief document as text: Markdown and other text as it is, Word, PDF and Excel through their readers.
    None for a file that holds no readable text (it is named in the trace and left out)."""
    kind = f.suffix.lower()
    try:
        if kind == ".docx":
            import docx
            d = docx.Document(str(f)); out = [p.text for p in d.paragraphs]
            for t in d.tables:
                out.append("")
                for row in t.rows: out.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(out)
        elif kind == ".pdf":
            import pypdf
            text = "\n\n".join((pg.extract_text() or "") for pg in pypdf.PdfReader(str(f)).pages)
        elif kind in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(str(f), data_only=True, read_only=True); out = []
            for ws in wb.worksheets:
                out.append(f"--- sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None and str(c).strip() for c in row): out.append(" | ".join("" if c is None else str(c) for c in row))
            text = "\n".join(out)
        else:
            raw = f.read_bytes()
            if b"\x00" in raw[:4096]: return None
            try: text = raw.decode("utf-8")
            except UnicodeDecodeError: text = raw.decode("latin-1")
    except Exception as e:
        return None
    return text if text.strip() else None

def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    opt = {a.split("=")[0]: (a.split("=", 1)[1] if "=" in a else True) for a in argv if a.startswith("--")}
    for i, a in enumerate(argv):                                   # "--slot NAME" as well as "--slot=NAME"
        if a in ("--slot", "--spec") and i + 1 < len(argv): opt[a] = argv[i + 1]; args.remove(argv[i + 1])
    briefs, api, url = inputs(args)
    if not briefs: sys.exit(__doc__)
    slot = opt.get("--slot") or pathlib.Path([a for a in args if a != url and pathlib.Path(a) != api][-1]).stem      # the last brief file, or the folder
    global TRACE
    trace = TRACE = llm.Trace(str(ROOT / "traces" / (slot + ".log")))
    trace.log(f"builder run for slot '{slot}' | briefs: {', '.join(str(b) for b in briefs)} | tool API: {api or 'none given, taken from the documents'}" + (f" | bank services at: {url}" if url else "") + f" | builder model: {llm.MODEL}")
    texts = [(b.name, read_text(b)) for b in briefs]
    for name, t in texts:
        if t is None: trace.log(f"brief file {name} holds no readable text and was left out")
    texts = [(n, t) for n, t in texts if t is not None]
    for i, (n, t) in enumerate(texts):
        if len(t) > LIMIT: trace.log(f"brief file {n} is {len(t)} characters long, the first {LIMIT} are read"); texts[i] = (n, t[:LIMIT])
    if not texts: sys.exit("no brief file could be read as text")
    trace.log("documents read: " + ", ".join(f"{n} ({len(t)} characters)" for n, t in texts))
    if api is not None: openapi = json.loads(api.read_text())
    elif opt.get("--spec") and (ROOT / "out" / (slot + ".openapi.json")).exists(): openapi = json.loads((ROOT / "out" / (slot + ".openapi.json")).read_text())
    else: openapi = llm.run(extract.tool_api(trace, texts))
    if url:
        # the tools call the bank's services at the address given on the command line; a path the documents put under the bank's own base
        # address (for example /api/v2) is kept, unless the address given already has a path
        import urllib.parse
        doc = urllib.parse.urlparse(str(((openapi.get("servers") or [{}])[0] or {}).get("url") or "")); given_url = urllib.parse.urlparse(url)
        base = url.rstrip("/") + (doc.path.rstrip("/") if given_url.path.strip("/") == "" else "")
        openapi["servers"] = [{"url": base}]; trace.log("the tools call the bank's services at " + base)
    api_text = json.dumps(openapi, indent=1)
    (ROOT / "out").mkdir(exist_ok=True); (ROOT / "out" / (slot + ".openapi.json")).write_text(api_text)
    (ROOT / "out").mkdir(exist_ok=True)
    # the engine handles one kind of call and needs services for four roles; for any other set of services the free-form way runs (freeform.py)
    from builder import freeform
    named = freeform.fits_engine(openapi)
    free = bool(opt.get("--freeform"))
    if opt.get("--spec"):
        spec = json.load(open(opt["--spec"])); trace.log("spec read from " + str(opt["--spec"]) + " (no model call)")
        free = free or "steps" in spec
    else:
        if free: trace.log("WAY: free-form (--freeform given)")
        else:
            # the engine first: its services may carry other names than its roles (tools.<role>.service). Only when the reading finds no service
            # for a role the engine cannot do without, or cannot be made to pass the check, the free-form way runs
            if not named:
                # other names than the engine's roles: one short question to the model settles whether these services are the engine's kind of call at all
                try: played = llm.run(freeform.roles_played(trace, openapi))
                except Exception as e: played = {}; trace.log("the question which services play the engine's roles got no usable answer: " + str(e)[:200])
                lacking = [r for r in freeform.ROLES if not played.get(r)]
                if lacking: free = True; trace.log("WAY: free-form, no service of the tool API plays the engine's role(s) " + ", ".join(lacking))
            if not free:
                trace.log("WAY: engine, " + ("the tool API has a service under the name of every role the engine needs" if named else
                          "the tool API names its services differently, the reading maps the engine's roles to them: " + json.dumps(played)))
                try: spec = llm.run(extract.extract(trace, texts, openapi, api_text))
                except Exception as e:                 # a reading that fails is not the end of the run: the free-form way still makes an agent
                    free = True; trace.log("WAY: free-form, the engine's reading did not fit these services: " + type(e).__name__ + ": " + " ".join(str(e).split())[:500])
        if free: spec = llm.run(freeform.read(trace, texts, openapi, api_text))
        (ROOT / "out" / (slot + ".spec.json")).write_text(json.dumps(spec, indent=1))
        trace.log(f"spec written to out/{slot}.spec.json")
    if not free:
        try: wf = engine.lint(engine.build_workflow(spec, openapi))
        except Exception as e:
            if opt.get("--spec"): raise
            free = True; trace.log("WAY: free-form, the engine could not compile this policy: " + type(e).__name__ + ": " + " ".join(str(e).split())[:500])
            spec = llm.run(freeform.read(trace, texts, openapi, api_text)); (ROOT / "out" / (slot + ".spec.json")).write_text(json.dumps(spec, indent=1))
    if free: wf = freeform.build_workflow(spec, openapi)
    (ROOT / "out" / (slot + ".workflow.json")).write_text(json.dumps(wf, indent=1))
    trace.log(f"workflow compiled: {len(wf['nodes'])} steps, {len(json.dumps(wf))} bytes, checks passed")
    if free: trace.log("SELF-TEST skipped: it tests the engine's in-call code, which the free-form workflow does not use")
    else:
        try: selftest(spec, openapi, trace)
        except Exception as e: trace.log("SELF-TEST against a fake bank FAILED (the agent is still deployed): " + type(e).__name__ + ": " + str(e)[:400])
    if opt.get("--no-deploy"): trace.end("not deployed"); return
    import freya
    from builder import deploy
    deploy.again("dry run", lambda: freya.validate_workflow(slot, wf), trace); trace.log("Freya dry run: valid")
    ids = deploy.save(freya, slot, wf, str(ROOT / ".freya-agents.json"), trace)
    trace.log(f"CREATED ON FREYA: workflow {ids['workflow_id']} agent {ids['agent_id']}" + (f" (updated in place, published version {ids['version']})" if ids.get("version") else " (new workflow)"))
    (ROOT / "submission").mkdir(exist_ok=True)
    out = deploy.again("exporting the workflow", lambda: freya.export_workflow(ids["workflow_id"], str(ROOT / "submission" / (slot + ".workflow.json"))), trace)
    trace.log("exported to " + os.path.relpath(out, ROOT))
    try:
        r = deploy.again("chat test", lambda: freya.api("POST", "/api/v2/chat/message", {"workflowId": ids["workflow_id"], "message": "Hello?", "customerVariables": spec.get("test_variables") or {}}), trace, tries=4)
        trace.log("TEST over the chat route, first line of the agent: " + json.dumps(r.get("response")))
        try: freya.api("POST", "/api/v2/chat/end", {"callId": r["callId"]})
        except Exception: pass
    except Exception as e:
        trace.log("TEST over the chat route did not answer: " + str(e)[:200])
    trace.end(f"workflow {ids['workflow_id']} agent {ids['agent_id']}")
    print(json.dumps(ids))

TRACE = None

if __name__ == "__main__":
    try: main(sys.argv[1:])
    except BaseException as e:                     # a run that breaks off still ends its trace, with the reason and what it cost
        if TRACE is not None and not TRACE.ended: TRACE.end("FAILED, nothing was delivered by this run: " + type(e).__name__ + ": " + " ".join(str(e).split())[:400])
        raise
